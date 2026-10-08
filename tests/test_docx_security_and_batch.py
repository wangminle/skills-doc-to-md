"""Issue #1 专项测试：恶意 DOCX 资源耗尽防线、SHA-256 完成标记、
批处理单文档超时与半成品清理。"""

import errno
import hashlib
import importlib.util
import io
import os
import re
import signal
import struct
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest import mock


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = WORKSPACE_ROOT / "skills" / "docx-to-markdown" / "scripts"
BATCH_SCRIPT = SCRIPTS_DIR / "batch_convert.py"
def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def build_minimal_docx(path, media_entries=None, extra_entries=None):
    """构造最小可用 DOCX（含 word/document.xml），可注入 media/其他条目。"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        zf.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body><w:p><w:r><w:t>测试正文</w:t></w:r></w:p></w:body></w:document>',
        )
        for name, data in (media_entries or {}).items():
            zf.writestr(f"word/media/{name}", data)
        for name, data in (extra_entries or {}).items():
            zf.writestr(name, data)


def rewrite_docx(path, comment=b"changed"):
    """重写 ZIP（仅改 comment）使文件字节变化但内容仍为有效 DOCX。"""
    with zipfile.ZipFile(path, "r") as zf:
        entries = [(info.filename, zf.read(info.filename)) for info in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.comment = comment
        for name, data in entries:
            zf.writestr(name, data)


class TestZipSecurityValidation(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def _make_docx(self, tmpdir, **kwargs):
        path = os.path.join(tmpdir, "input.docx")
        build_minimal_docx(path, **kwargs)
        return path

    def test_security_error_is_value_error_subclass(self):
        # 继承 ValueError 兼容既有处理，同时可被调用方精确区分
        self.assertTrue(issubclass(self.convert.DocxSecurityError, ValueError))

    def test_real_docx_passes_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._make_docx(tmp)
            with zipfile.ZipFile(path, "r") as zf:
                self.convert.validate_docx_zip_security(zf)  # 不抛异常即通过

    def _assert_rejected(self, tmpdir, patched_limits, **kwargs):
        path = self._make_docx(tmpdir, **kwargs)
        with zipfile.ZipFile(path, "r") as zf:
            with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, patched_limits):
                with self.assertRaises(self.convert.DocxSecurityError):
                    self.convert.validate_docx_zip_security(zf)

    def test_total_uncompressed_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._assert_rejected(
                tmp,
                {"total_uncompressed": 1024 * 1024},
                extra_entries={"big.bin": b"\0" * (2 * 1024 * 1024)},
            )

    def test_entry_uncompressed_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._assert_rejected(
                tmp,
                {"entry_uncompressed": 512 * 1024},
                extra_entries={"big.bin": b"\0" * (1024 * 1024)},
            )

    def test_entry_compression_ratio_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            # 100KB 高度可压缩数据，压缩后远小于 1/10
            self._assert_rejected(
                tmp,
                {"entry_ratio": 10},
                extra_entries={"sponge.bin": b"A" * (100 * 1024)},
            )

    def test_total_compression_ratio_limit_with_size_floor(self):
        with tempfile.TemporaryDirectory() as tmp:
            # 压缩后总大小低于门槛时不判定（避免小文件舍入误伤）……
            path = self._make_docx(tmp, extra_entries={"sponge.bin": b"A" * (1024 * 1024)})
            with zipfile.ZipFile(path, "r") as zf:
                with mock.patch.dict(
                    self.convert.DOCX_SECURITY_LIMITS,
                    {"entry_ratio": 10 ** 6, "total_ratio": 10,
                     "total_ratio_min_compressed": 10 * 1024 * 1024},
                ):
                    self.convert.validate_docx_zip_security(zf)  # 通过
            # ……超过门槛且总压缩比超限时拒绝
            with zipfile.ZipFile(path, "r") as zf:
                with mock.patch.dict(
                    self.convert.DOCX_SECURITY_LIMITS,
                    {"entry_ratio": 10 ** 6, "total_ratio": 10,
                     "total_ratio_min_compressed": 512},
                ):
                    with self.assertRaises(self.convert.DocxSecurityError):
                        self.convert.validate_docx_zip_security(zf)

    def test_image_count_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._assert_rejected(
                tmp,
                {"image_count": 2},
                media_entries={"image1.png": b"\x89PNG\r\n\x1a\n",
                               "image2.png": b"\x89PNG\r\n\x1a\n",
                               "image3.png": b"\x89PNG\r\n\x1a\n"},
            )

    def test_image_file_size_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._assert_rejected(
                tmp,
                {"image_file_size": 1024 * 1024},
                media_entries={"image1.png": b"\x89PNG\r\n\x1a\n" + b"\0" * (2 * 1024 * 1024)},
            )

    def test_embedded_excel_size_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            self._assert_rejected(
                tmp,
                {"embedded_excel_size": 1024 * 1024},
                extra_entries={"word/embeddings/sheet1.xlsx": b"\0" * (2 * 1024 * 1024)},
            )

    def test_conversion_rejects_zip_bomb_before_any_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._make_docx(tmp, extra_entries={"sponge.bin": b"A" * (10 * 1024 * 1024)})
            with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"entry_ratio": 10}):
                with self.assertRaises(self.convert.DocxSecurityError):
                    self.convert.convert_docx_to_markdown(path, os.path.join(tmp, "out"))
            # 安全拒绝发生在任何输出创建之前
            self.assertFalse(os.path.exists(os.path.join(tmp, "out")))


class TestImagePixelCount(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_png(self):
        ihdr = struct.pack(">II", 20000, 15000)
        data = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + ihdr + b"\x08\x06\x00\x00\x00"
        self.assertEqual(self.convert.image_pixel_count(data), 20000 * 15000)

    def test_jpeg(self):
        data = (b"\xff\xd8" + b"\xff\xc0" + b"\x00\x11" + b"\x08"
                + struct.pack(">HH", 3000, 4000)  # height, width
                + b"\x03" + b"\x01\x11\x00" * 3 + b"\xff\xd9")
        self.assertEqual(self.convert.image_pixel_count(data), 3000 * 4000)

    def test_gif(self):
        data = b"GIF89a" + struct.pack("<HH", 1000, 2000)
        self.assertEqual(self.convert.image_pixel_count(data), 1000 * 2000)

    def test_bmp(self):
        data = b"BM" + b"\x00" * 16 + struct.pack("<ii", 1000, 2000) + b"\x00" * 8
        self.assertEqual(self.convert.image_pixel_count(data), 1000 * 2000)

    def test_webp_vp8x(self):
        buf = bytearray(34)
        buf[0:4] = b"RIFF"
        buf[8:12] = b"WEBP"
        buf[12:16] = b"VP8X"
        buf[24:27] = (2000 - 1).to_bytes(3, "little")
        buf[27:30] = (3000 - 1).to_bytes(3, "little")
        self.assertEqual(self.convert.image_pixel_count(bytes(buf)), 2000 * 3000)

    def test_tiff(self):
        entries = struct.pack("<H", 2)
        entries += struct.pack("<HHII", 256, 4, 1, 2000)  # ImageWidth
        entries += struct.pack("<HHII", 257, 4, 1, 1500)  # ImageLength
        entries += struct.pack("<I", 0)
        data = b"II*\x00" + struct.pack("<I", 8) + entries
        self.assertEqual(self.convert.image_pixel_count(data), 2000 * 1500)

    def test_unknown_format_returns_none(self):
        self.assertIsNone(self.convert.image_pixel_count(b"\xd7\xcd\xc6\x9a" + b"\x00" * 20))

    def test_pixel_bomb_rejected_during_extraction(self):
        # 头部声明 20000x15000（3 亿像素）的 PNG，默认 5000 万上限即拒绝
        ihdr = struct.pack(">II", 20000, 15000)
        png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + ihdr + b"\x08\x06\x00\x00\x00"
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bomb.docx")
            build_minimal_docx(path, media_entries={"image1.png": png})
            with self.assertRaises(self.convert.DocxSecurityError):
                self.convert.convert_docx_to_markdown(path, os.path.join(tmp, "out"))


class TestSanitizeStem(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_plain_name_unchanged(self):
        self.assertEqual(self.convert.sanitize_stem("设备预约V2.7.2"), "设备预约V2.7.2")

    def test_nfkc_only_normalizes_without_hash_suffix(self):
        # 全角→半角归一化过于普遍（尤其中文文档），不附加 hash；
        # 由此产生的罕见碰撞由 sentinel 源哈希校验兜底
        self.assertEqual(self.convert.sanitize_stem("自研语义VAD（测试）"), "自研语义VAD(测试)")

    def test_forbidden_char_replacement_appends_distinct_hash(self):
        # a:b 与 a_b 均清洗为 a_b，但 hash 后缀不同，不共享输出目录
        self.assertNotEqual(self.convert.sanitize_stem("a:b"), self.convert.sanitize_stem("a_b"))
        self.assertTrue(self.convert.sanitize_stem("a:b").startswith("a_b_"))

    def test_quote_removal_appends_hash(self):
        stemmed = self.convert.sanitize_stem('文档"引用"')
        self.assertTrue(stemmed.startswith("文档引用_"))

    def test_long_name_truncated_with_hash(self):
        long_name = "很长的文档名" * 30
        stemmed = self.convert.sanitize_stem(long_name)
        self.assertLessEqual(len(stemmed), 120)
        self.assertIn("_", stemmed)


class TestConversionSentinel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_conversion_writes_json_sentinel_with_source_hash(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = os.path.join(tmp, "input.docx")
            build_minimal_docx(docx_path)
            md_path = self.convert.convert_docx_to_markdown(docx_path, tmp)
            out_dir = os.path.dirname(md_path)

            sentinel = self.convert.read_conversion_sentinel(out_dir)
            self.assertIsNotNone(sentinel)
            self.assertEqual(sentinel["source_sha256"], self.convert.sha256_file(docx_path))
            self.assertEqual(sentinel["folder_name"], os.path.basename(out_dir))
            self.assertEqual(sentinel["on_limit"], "reject")

    def test_v016_json_sentinel_defaults_to_reject_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, ".converted").write_text(
                '{"folder_name": "x", "source_sha256": "abc"}', encoding="utf-8"
            )
            sentinel = self.convert.read_conversion_sentinel(tmp)
            self.assertEqual(sentinel["on_limit"], "reject")

    def test_unknown_sentinel_policy_is_invalid(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, ".converted").write_text(
                '{"folder_name": "x", "source_sha256": "abc", "on_limit": "ignore"}',
                encoding="utf-8",
            )
            self.assertIsNone(self.convert.read_conversion_sentinel(tmp))

    def test_legacy_plain_text_sentinel_is_invalid(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, ".converted").write_text("done", encoding="utf-8")
            self.assertIsNone(self.convert.read_conversion_sentinel(tmp))

    def test_corrupt_or_incomplete_sentinel_is_invalid(self):
        for payload in ('{"folder_name": "x"}',  # 缺 source_sha256
                        "not json at all",
                        '["list", "not", "dict"]'):
            with tempfile.TemporaryDirectory() as tmp:
                Path(tmp, ".converted").write_text(payload, encoding="utf-8")
                self.assertIsNone(self.convert.read_conversion_sentinel(tmp))

    def test_missing_sentinel_is_invalid(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertIsNone(self.convert.read_conversion_sentinel(tmp))


class TestBatchConvert(unittest.TestCase):
    def setUp(self):
        self.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        self.batch = load_module("batch_convert_module", SCRIPTS_DIR / "batch_convert.py")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.src = os.path.join(self.tmp.name, "src")
        self.out = os.path.join(self.tmp.name, "out")
        os.makedirs(self.src)
        self.docx_path = os.path.join(self.src, "input.docx")
        build_minimal_docx(self.docx_path)
        base_name = os.path.splitext(os.path.basename(self.docx_path))[0]
        self.folder_name = self.convert.sanitize_stem(base_name)
        self.target_dir = os.path.join(self.out, self.folder_name)

    def _run_batch(self, **kwargs):
        with self.assertLogs("batch_convert_module", level="INFO") as cm:
            self.batch.batch_convert(self.src, self.out, **kwargs)
        return [line for line in cm.output]

    def _run_batch_summary(self, src=None, **kwargs):
        with self.assertLogs("batch_convert_module", level="INFO"):
            return self.batch.batch_convert(src or self.src, self.out, **kwargs)

    def test_skip_when_complete_and_source_unchanged(self):
        self._run_batch()
        self.assertTrue(os.path.isfile(os.path.join(self.target_dir, f"{self.folder_name}.md")))

        logs = self._run_batch()
        self.assertTrue(any("源文件未变更，跳过" in line for line in logs))
        self.assertTrue(any("成功 0 个, 跳过 1 个, 失败 0 个" in line for line in logs))

    def test_source_change_triggers_reconversion_without_force(self):
        self._run_batch()
        rewrite_docx(self.docx_path)  # 字节变化，内容仍为有效 DOCX

        logs = self._run_batch()
        self.assertTrue(any("重新转换" in line for line in logs))
        sentinel = self.convert.read_conversion_sentinel(self.target_dir)
        self.assertEqual(sentinel["source_sha256"], self.convert.sha256_file(self.docx_path))

    def test_legacy_sentinel_output_is_reconverted(self):
        self._run_batch()
        Path(self.target_dir, ".converted").write_text("done", encoding="utf-8")  # 旧格式

        logs = self._run_batch()
        self.assertFalse(any("源文件未变更，跳过" in line for line in logs))
        self.assertIsNotNone(self.convert.read_conversion_sentinel(self.target_dir))

    def test_tampered_sentinel_hash_is_reconverted(self):
        self._run_batch()
        sentinel_path = Path(self.target_dir, ".converted")
        sentinel_path.write_text(
            '{"folder_name": "%s", "source_sha256": "%s"}' % (self.folder_name, "0" * 64),
            encoding="utf-8",
        )

        logs = self._run_batch()
        self.assertFalse(any("源文件未变更，跳过" in line for line in logs))

    def test_failed_conversion_cleans_half_finished_output(self):
        os.makedirs(self.target_dir)
        Path(self.target_dir, "half.md").write_text("半成品", encoding="utf-8")

        with mock.patch.object(
            self.batch, "convert_docx_to_markdown", side_effect=RuntimeError("boom")
        ):
            logs = self._run_batch()

        self.assertTrue(any("失败: boom" in line for line in logs))
        self.assertFalse(os.path.exists(self.target_dir))  # 半成品被清理

    def test_security_rejection_counts_as_failure_and_cleans_output(self):
        os.makedirs(self.target_dir)
        # 注意用 batch 命名空间中的异常类：batch 模块经 sys.path 导入的
        # convert_docx 与本测试 spec 加载的是不同实例，isinstance 不互通
        with mock.patch.object(
            self.batch,
            "convert_docx_to_markdown",
            side_effect=self.batch.DocxSecurityError("zip bomb"),
        ):
            logs = self._run_batch()

        self.assertTrue(any("安全拒绝" in line for line in logs))
        self.assertTrue(any("失败 1 个" in line for line in logs))
        self.assertFalse(os.path.exists(self.target_dir))

    def test_timeout_kills_hanging_conversion_and_cleans_output(self):
        if not hasattr(signal, "SIGALRM"):
            self.skipTest("平台无 SIGALRM（Windows），超时路径自动跳过")

        def slow_convert(*args, **kwargs):
            time.sleep(30)
            return "never"

        os.makedirs(self.target_dir)
        with mock.patch.object(self.batch, "convert_docx_to_markdown", side_effect=slow_convert):
            logs = self._run_batch(timeout=1)

        self.assertTrue(any("超时" in line for line in logs))
        self.assertFalse(os.path.exists(self.target_dir))

    def test_run_with_timeout_restores_previous_handler(self):
        if not hasattr(signal, "SIGALRM"):
            self.skipTest("平台无 SIGALRM")

        old_handler = signal.getsignal(signal.SIGALRM)
        try:
            self.assertEqual(self.batch._run_with_timeout(lambda: "ok", 30), "ok")
            self.assertEqual(signal.getsignal(signal.SIGALRM), old_handler)
            self.assertEqual(signal.alarm(0), 0)  # 无未决 alarm
        finally:
            signal.signal(signal.SIGALRM, old_handler)

    def test_run_with_timeout_disabled_for_non_positive(self):
        self.assertEqual(self.batch._run_with_timeout(lambda: "ok", 0), "ok")
        self.assertEqual(self.batch._run_with_timeout(lambda: "ok", -1), "ok")

    def test_nfkc_normalized_name_collision_keeps_both_outputs(self):
        # BUG-014 回归：A 与全角Ａ经 NFKC 归一化后同名，不得共用输出目录
        # 互相覆盖（sentinel 只能检出哈希不一致触发重转，防不了批内覆盖）
        build_minimal_docx(os.path.join(self.src, "A.docx"))
        fullwidth = os.path.join(self.src, "Ａ.docx")
        build_minimal_docx(fullwidth)
        rewrite_docx(fullwidth, comment=b"fullwidth")  # 与 A.docx 字节不同

        summary = self._run_batch_summary()

        # setUp 的 input.docx 正常输出；A 与 Ａ 必须各占一个独立目录
        dirs = sorted(os.listdir(self.out))
        self.assertEqual(len(dirs), 3, f"应有三个独立输出目录，实际 {dirs}")
        self.assertIn("A", dirs)
        other = [d for d in dirs if d not in ("A", "input")][0]
        self.assertTrue(other.startswith("A_"), f"消歧名应保留原名前缀: {other}")
        for d in dirs:
            self.assertTrue(Path(self.out, d, f"{d}.md").is_file())
            self.assertIsNotNone(self.convert.read_conversion_sentinel(
                os.path.join(self.out, d)))
        self.assertEqual(summary, {"success": 3, "skipped": 0, "failed": 0})

        # 消歧只依赖文件名，第二轮分配结果一致，全部输出命中 sentinel 跳过
        logs = self._run_batch()
        self.assertTrue(any("成功 0 个, 跳过 3 个, 失败 0 个" in line for line in logs))

    def test_second_order_name_collision_keeps_all_three_outputs(self):
        # BUG-018 回归：hash 消歧名恰与批内自然文件名相同时，最终分配名
        # （含 _2 序号）必须透传给转换器，否则转换器仍写回旧目录覆盖结果
        digest = hashlib.sha256("Ａ".encode("utf-8")).hexdigest()[:8]
        build_minimal_docx(os.path.join(self.src, "A.docx"))
        build_minimal_docx(os.path.join(self.src, f"A_{digest}.docx"))
        fullwidth = os.path.join(self.src, "Ａ.docx")
        build_minimal_docx(fullwidth)
        rewrite_docx(fullwidth, comment=b"fullwidth")

        summary = self._run_batch_summary()

        # setUp 的 input.docx 正常输出；三个同名文件各占一个独立目录
        dirs = sorted(os.listdir(self.out))
        self.assertEqual(dirs, sorted(["A", f"A_{digest}", f"A_{digest}_2", "input"]))
        for d in dirs:
            self.assertTrue(Path(self.out, d, f"{d}.md").is_file(), f"缺少 {d}.md")
            sentinel = self.convert.read_conversion_sentinel(os.path.join(self.out, d))
            self.assertEqual(sentinel["folder_name"], d)
        self.assertEqual(summary, {"success": 4, "skipped": 0, "failed": 0})

        # 分配跨批次稳定，第二轮全部命中 sentinel 跳过
        logs = self._run_batch()
        self.assertTrue(any("成功 0 个, 跳过 4 个, 失败 0 个" in line for line in logs))

    def test_batch_returns_failure_summary(self):
        # BUG-015 回归（API 侧）：失败数必须反映在返回值中供调用方判定
        summary_ok = self._run_batch_summary()
        self.assertEqual(summary_ok, {"success": 1, "skipped": 0, "failed": 0})

        # force 绕过首轮已写好的 sentinel，否则会走跳过而不是失败分支
        with mock.patch.object(
            self.batch, "convert_docx_to_markdown", side_effect=RuntimeError("boom")
        ):
            summary_fail = self._run_batch_summary(force=True)
        self.assertEqual(summary_fail, {"success": 0, "skipped": 0, "failed": 1})

        # 空目录：告警但不算失败
        empty_src = os.path.join(self.tmp.name, "empty")
        os.makedirs(empty_src)
        summary_empty = self._run_batch_summary(src=empty_src)
        self.assertEqual(summary_empty, {"success": 0, "skipped": 0, "failed": 0})

    def test_batch_cli_exit_code_nonzero_on_failure(self):
        # BUG-015 回归（CLI 侧）：任一文档失败退出码 1，全部跳过/成功退出码 0
        Path(self.src, "broken.docx").write_bytes(b"not a zip at all")
        result = subprocess.run(
            [sys.executable, str(BATCH_SCRIPT), self.src, self.out],
            capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        # 失败文档不拖累同批正常文档
        self.assertTrue(Path(self.target_dir, f"{self.folder_name}.md").is_file())

        os.remove(os.path.join(self.src, "broken.docx"))
        result_ok = subprocess.run(
            [sys.executable, str(BATCH_SCRIPT), self.src, self.out],
            capture_output=True, text=True)
        self.assertEqual(result_ok.returncode, 0)

    def test_mixed_case_docx_extension_discovered(self):
        # BUG-016 回归：.Docx/.DOCX/.doCx 等混合大小写扩展名都要进入批处理，
        # 同名的目录与非 docx 文件仍被忽略
        os.remove(self.docx_path)
        build_minimal_docx(os.path.join(self.src, "Mixed.Docx"))
        build_minimal_docx(os.path.join(self.src, "b.DOCX"))
        build_minimal_docx(os.path.join(self.src, "c.doCx"))
        Path(self.src, "note.txt").write_text("hello", encoding="utf-8")
        Path(self.src, "archive.doc").write_bytes(b"old format")
        os.makedirs(os.path.join(self.src, "adir.docx"))

        summary = self._run_batch_summary()

        self.assertEqual(summary, {"success": 3, "skipped": 0, "failed": 0})
        for name in ("Mixed", "b", "c"):
            self.assertTrue(
                Path(self.out, name, f"{name}.md").is_file(), f"缺少 {name} 的输出")


class TestMediaDirectoryEntry(unittest.TestCase):
    """BUG-017 回归：显式 word/media/ 目录 entry 不是图片。"""

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def _build_docx(self, tmp):
        path = os.path.join(tmp, "dir_entry.docx")
        build_minimal_docx(
            path,
            media_entries={"image1.png": b"\x89PNG\r\n\x1a\n" + b"\0" * 32},
            extra_entries={"word/media/": b""},  # 显式目录 entry
        )
        return path

    def test_directory_entry_not_extracted_as_image(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._build_docx(tmp)
            with zipfile.ZipFile(path, "r") as zf:
                self.assertTrue(any(info.is_dir() for info in zf.infolist()))

            assets_dir = os.path.join(tmp, "assets")
            os.makedirs(assets_dir)
            image_by_hash, _, _ = self.convert.extract_content_from_docx(
                path, assets_dir)

            # 不产生空内容的 assets/.png 伪文件与空数据 hash 映射
            self.assertEqual(os.listdir(assets_dir), ["image1.png"])
            empty_sha = hashlib.sha256(b"").hexdigest()
            self.assertNotIn(empty_sha, image_by_hash)

    def test_directory_entry_does_not_consume_skip_quota(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = self._build_docx(tmp)
            assets_dir = os.path.join(tmp, "assets")
            os.makedirs(assets_dir)
            skip_state = {}

            # 配额恰好 1 张：目录 entry 不得挤占唯一名额
            with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 1}):
                images, _, _ = self.convert.extract_content_from_docx(
                    path, assets_dir, on_limit="skip", skip_state=skip_state)

            self.assertEqual(len(images), 1)
            self.assertEqual(skip_state["skipped"], [])


class TestDotSegmentImageQuota(unittest.TestCase):
    """BUG-038 回归：含 . / .. 路径段的关系目标不得绕过图片数量配额。

    Mammoth 的 uri_to_zip_entry_name 只按字面拼接条目名（/ 开头取
    uri[1:]，否则 "word/" + uri），不做 normpath；扫描器此前只按规范化名
    查条目，word/custom/../imageN.png 这类字面条目两边对不上，配额失效。
    """

    _DOC_NS = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
    )
    _IMG_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

    # (关系 Target 写法, ZIP 内字面条目名)：Mammoth 能读到字面条目，
    # 但规范化后与条目名不一致——正是绕过配额的构造。
    _VARIANTS = {
        "dotdot": ("custom/../image{i}.png", "word/custom/../image{i}.png"),
        "dot": ("./image{i}.png", "word/./image{i}.png"),
        "root-dot": ("/custom/./image{i}.png", "custom/./image{i}.png"),
    }

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        # 三张字节互不相同的图：配额被绕过时会全部落盘，红证据直观
        cls.image_data = {
            i: (b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR"
                + struct.pack(">II", i + 1, i + 1) + b"\x08\x06\x00\x00\x00")
            for i in (1, 2, 3)
        }

    def _build_docx(self, path, target_fmt, entry_fmt):
        drawing = (
            '<w:drawing><wp:inline><wp:extent cx="100" cy="100"/><a:graphic><a:graphicData '
            'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr>'
            '<pic:cNvPr id="{i}" name="i{i}"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
            '<a:blip r:embed="rId{i}"/></pic:blipFill><pic:spPr/></pic:pic></a:graphicData>'
            '</a:graphic></wp:inline></w:drawing>'
        )
        document = (
            f'<?xml version="1.0"?><w:document {self._DOC_NS}><w:body>'
            + "".join(f'<w:p><w:r>{drawing.format(i=i)}</w:r></w:p>' for i in (1, 2, 3))
            + '<w:p><w:r><w:t>正文</w:t></w:r></w:p></w:body></w:document>'
        )
        rels = (
            '<?xml version="1.0"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(
                f'<Relationship Id="rId{i}" Type="{self._IMG_TYPE}" '
                f'Target="{target_fmt.format(i=i)}"/>' for i in (1, 2, 3))
            + "</Relationships>"
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
            zf.writestr("word/document.xml", document)
            zf.writestr("word/_rels/document.xml.rels", rels)
            for i in (1, 2, 3):
                zf.writestr(entry_fmt.format(i=i), self.image_data[i])

    def _variants_tmp(self):
        for name, (target_fmt, entry_fmt) in self._VARIANTS.items():
            tmp = tempfile.TemporaryDirectory()
            self.addCleanup(tmp.cleanup)
            path = os.path.join(tmp.name, f"{name}.docx")
            self._build_docx(path, target_fmt, entry_fmt)
            yield name, path, os.path.join(tmp.name, "out")

    def test_scanner_counts_literal_dot_segment_entries(self):
        for name, path, _ in self._variants_tmp():
            with zipfile.ZipFile(path, "r") as zf:
                counted = self.convert._document_image_part_names(zf)
            # 三张字面条目全部计入，且键就是 ZIP 里真实存在的条目名
            self.assertEqual(len(counted), 3, (name, sorted(counted)))
            entry_fmt = self._VARIANTS[name][1]
            for i in (1, 2, 3):
                self.assertIn(entry_fmt.format(i=i), counted, name)

    def test_reject_mode_enforces_count_quota(self):
        for name, path, out in self._variants_tmp():
            with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 1}):
                with self.assertRaises(self.convert.DocxSecurityError, msg=name):
                    self.convert.convert_docx_to_markdown(path, out)
            self.assertFalse(os.path.exists(out), name)

    def test_skip_mode_quota_leaves_single_image(self):
        for name, path, out in self._variants_tmp():
            with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 1}):
                md_path = self.convert.convert_docx_to_markdown(
                    path, out, on_limit="skip")
            assets = os.path.join(os.path.dirname(md_path), "assets")
            self.assertEqual(os.listdir(assets), ["image1.png"], name)
            content = Path(md_path).read_text(encoding="utf-8")
            self.assertEqual(
                len(re.findall(r"!\[\]\(assets/[^)]+\)", content)), 1, name)
            self.assertEqual(content.count("【图片已跳过：图片数量超过上限】"), 2, name)


class TestDirNamedEntryDefense(unittest.TestCase):
    """BUG-041/043 回归：名字以 / 结尾但带内容的条目不得绕过防线。

    zipfile 对 is_dir() 条目（名字以 / 结尾）仍能按字面名读取真实数据：
    声明非零内容的目录名条目必须拒绝（单文件/压缩比/总量三道检查之前
    不得跳过）；文件名形式的目录名媒体条目（word/media/img.png/）照常
    计入图片配额扫描。纯目录占位条目（word/media/ 本身）不是图片。
    """

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_dir_named_entry_with_declared_content_rejected(self):
        # 修复前：该条目在体积/压缩比/总量统计之前被 is_dir 跳过，
        # 288KB 输入即可让 mammoth 解压解析 120MB XML（BUG-041 红证据）
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "dirbomb.docx")
            build_minimal_docx(
                path, extra_entries={"word/real.xml/": b"A" * (2 * 1024 * 1024)})
            with zipfile.ZipFile(path, "r") as zf:
                self.assertTrue(zf.getinfo("word/real.xml/").is_dir())
                with self.assertRaisesRegex(
                        self.convert.DocxSecurityError, "目录名条目声明非零内容"):
                    self.convert.validate_docx_zip_security(zf)

    def test_empty_dir_placeholder_still_allowed(self):
        # 正常空目录占位条目（file_size 恒为 0）不受影响
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "plain.docx")
            build_minimal_docx(path, extra_entries={"word/media/": b""})
            with zipfile.ZipFile(path, "r") as zf:
                self.convert.validate_docx_zip_security(zf)  # 不抛即通过

    def test_dir_named_media_entries_counted_by_scan(self):
        # 修复前 _document_image_part_names 对 is_dir 条目一律跳过，扫描
        # 结果为空集（BUG-043 红证据：mammoth 仍能按字面名读取并落盘）
        png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(">II", 2, 2)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "dirname.docx")
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
                zf.writestr("word/document.xml", "<w:document/>")
                zf.writestr(
                    "word/_rels/document.xml.rels",
                    '<?xml version="1.0"?><Relationships '
                    'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    + "".join(
                        '<Relationship Id="rId%d" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                        'Target="media/img%d.png/"/>' % (i, i) for i in (1, 2, 3))
                    + "</Relationships>")
                for i in (1, 2, 3):
                    zf.writestr("word/media/img%d.png/" % i, png)
            with zipfile.ZipFile(path, "r") as zf:
                counted = self.convert._document_image_part_names(zf)
            self.assertEqual(
                counted,
                {"word/media/img1.png/", "word/media/img2.png/", "word/media/img3.png/"})

    def test_pure_dir_placeholder_name_not_counted(self):
        # word/media/ 本身（前缀后无文件名部分）不是图片数据，防止空名
        # 伪文件 assets/.png 与空数据 hash 映射
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "placeholder.docx")
            build_minimal_docx(path, extra_entries={"word/media/": b""})
            with zipfile.ZipFile(path, "r") as zf:
                counted = self.convert._document_image_part_names(zf)
            self.assertEqual(counted, set())


class TestRedirectedMainPartQuota(unittest.TestCase):
    """BUG-042 回归：正文 part 经 _rels/.rels 重定向后图片配额仍生效。

    mammoth 从 _rels/.rels 的 officeDocument 关系定位正文 part；扫描器
    此前只读硬编码的 word/document.xml(.rels)，正文 part 被重定向后
    reject 模式数量防线不生效、skip 模式回调兜底无配额写盘。
    """

    _DOC_NS = TestDotSegmentImageQuota._DOC_NS
    _IMG_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def _build_docx(self, path):
        drawing = (
            '<w:drawing><wp:inline><wp:extent cx="100" cy="100"/><a:graphic><a:graphicData '
            'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr>'
            '<pic:cNvPr id="{i}" name="i{i}"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
            '<a:blip r:embed="rId{i}"/></pic:blipFill><pic:spPr/></pic:pic></a:graphicData>'
            '</a:graphic></wp:inline></w:drawing>'
        )
        body = ("".join(f'<w:p><w:r>{drawing.format(i=i)}</w:r></w:p>' for i in (1, 2, 3))
                + '<w:p><w:r><w:t>正文</w:t></w:r></w:p>')
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
            # 诱饵：扫描器旧口径读的 part（无图片关系）
            zf.writestr("word/document.xml", "<w:document/>")
            zf.writestr("word/_rels/document.xml.rels",
                        '<?xml version="1.0"?><Relationships '
                        'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
            # 包关系把正文 part 指向 word/real.xml（mammoth 实际读取）
            zf.writestr(
                "_rels/.rels",
                '<?xml version="1.0"?><Relationships '
                'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                'Target="word/real.xml"/></Relationships>')
            zf.writestr(
                "word/real.xml",
                '<?xml version="1.0"?>'
                f'<w:document {self._DOC_NS}><w:body>{body}</w:body></w:document>')
            zf.writestr(
                "word/_rels/real.xml.rels",
                '<?xml version="1.0"?><Relationships '
                'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + "".join(
                    '<Relationship Id="rId%d" Type="%s" Target="custom/img%d.png"/>'
                    % (i, self._IMG_TYPE, i) for i in (1, 2, 3))
                + "</Relationships>")
            for i in (1, 2, 3):
                zf.writestr("word/custom/img%d.png" % i,
                            b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR"
                            + struct.pack(">II", i + 1, i + 1))

    def _docx(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = os.path.join(tmp.name, "redirect.docx")
        self._build_docx(path)
        return path, os.path.join(tmp.name, "out")

    def test_scan_resolves_redirected_main_part(self):
        # 修复前扫描只读硬编码 word/document.xml(.rels)：正文 part 与
        # 图片全部不可见（BUG-042 红证据）
        path, _ = self._docx()
        with zipfile.ZipFile(path, "r") as zf:
            parts = self.convert._mammoth_content_parts(zf)
            counted = self.convert._document_image_part_names(zf)
        self.assertEqual(parts[0], ("word/real.xml", "word/_rels/real.xml.rels"))
        self.assertEqual(counted, {f"word/custom/img{i}.png" for i in (1, 2, 3)})

    def test_reject_mode_enforces_quota_after_redirect(self):
        # 修复前 reject 模式（真实上限 500 时 501 张）全部落盘
        path, out = self._docx()
        with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 2}):
            with self.assertRaisesRegex(self.convert.DocxSecurityError, "图片数量超过上限"):
                self.convert.convert_docx_to_markdown(path, out)
        self.assertFalse(os.path.exists(out))

    def test_skip_mode_truncates_to_quota_after_redirect(self):
        # 修复前 skip 模式经回调兜底无配额写盘：501 张全部落盘
        path, out = self._docx()
        with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 2}):
            md_path = self.convert.convert_docx_to_markdown(path, out, on_limit="skip")
        assets = os.path.join(os.path.dirname(md_path), "assets")
        self.assertEqual(len(os.listdir(assets)), 2)
        content = Path(md_path).read_text(encoding="utf-8")
        self.assertEqual(content.count("【图片已跳过：图片数量超过上限】"), 1)

    def test_callback_fallback_budget_fail_closed(self):
        # 扫描口径与 mammoth 实际读取出现偏差（兜底分支写盘）时配额仍
        # 生效：mock 扫描结果为空集模拟扫描遗漏，预算耗尽即拒绝
        path, out = self._docx()
        with mock.patch.dict(self.convert.DOCX_SECURITY_LIMITS, {"image_count": 1}), \
                mock.patch.object(self.convert, "_document_image_part_names",
                                  return_value=set()):
            with self.assertRaisesRegex(self.convert.ResourceLimitExceeded,
                                        "mammoth 回调兜底"):
                self.convert.convert_docx_to_markdown(path, out)


class TestAssetWriteExclusiveOccupancy(unittest.TestCase):
    """BUG-047 回归：独占写入遇占用者时区分安全复用与竞争失败。

    分配与写入之间的窄竞争窗口内被第三方占用时，此前静默 return 使
    Markdown 引用指向符号链接/外来文件；现在同内容普通文件视为复用，
    符号链接/目录/不同内容的占用换下一候选重试。
    """

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_write_exclusive_reuses_identical_regular_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "img.png")
            data = b"\x89PNG\r\n\x1a\n"
            with open(path, "wb") as f:
                f.write(data)
            self.assertFalse(self.convert._write_asset_file_exclusive(path, data))

    def test_write_exclusive_rejects_foreign_occupants(self):
        data = b"\x89PNG\r\n\x1a\n"

        def write_other(p):
            with open(p, "wb") as f:
                f.write(b"OTHER")

        for kind, setup in (
            ("不同内容普通文件", write_other),
            ("目录", os.mkdir),
        ):
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "img.png")
                setup(path)
                with self.assertRaises(FileExistsError, msg=kind):
                    self.convert._write_asset_file_exclusive(path, data)
        with tempfile.TemporaryDirectory() as tmp:
            external = os.path.join(tmp, "external.bin")
            with open(external, "wb") as f:
                f.write(b"EXTERNAL")
            path = os.path.join(tmp, "img.png")
            os.symlink(external, path)
            with self.assertRaisesRegex(FileExistsError, ""):
                self.convert._write_asset_file_exclusive(path, data)

    def test_single_race_retries_with_next_candidate(self):
        # 一次性竞争：分配后、写入前候选名被符号链接抢占。修复前静默
        # 返回，Markdown 引用指向符号链接（红证据）；修复后换候选重写
        convert = self.convert
        png = (b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR"
               + struct.pack(">II", 2, 2) + b"\x08\x06\x00\x00\x00")
        with tempfile.TemporaryDirectory() as tmp:
            docx = os.path.join(tmp, "race.docx")
            with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
                zf.writestr(
                    "word/document.xml",
                    '<?xml version="1.0"?><w:document '
                    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
                    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                    '<w:body><w:p><w:r><w:drawing><wp:inline><a:graphic><a:graphicData '
                    'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
                    '<pic:nvPicPr><pic:cNvPr id="1" name="i"/><pic:cNvPicPr/></pic:nvPicPr>'
                    '<pic:blipFill><a:blip r:embed="rId1"/></pic:blipFill><pic:spPr/></pic:pic>'
                    '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
                    '<w:p><w:r><w:t>正文</w:t></w:r></w:p></w:body></w:document>')
                zf.writestr("word/_rels/document.xml.rels",
                            '<?xml version="1.0"?><Relationships '
                            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                            '<Relationship Id="rId1" '
                            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                            'Target="media/image1.png"/></Relationships>')
                zf.writestr("word/media/image1.png", png)

            external = os.path.join(tmp, "EXTERNAL_SECRET.bin")
            with open(external, "wb") as f:
                f.write(b"EXTERNAL-CONTENT-NOT-THE-IMAGE")

            out = os.path.join(tmp, "out")
            real_alloc = convert._allocate_asset_path
            races = []

            def racing_alloc(*args, **kwargs):
                name, path = real_alloc(*args, **kwargs)
                races.append(name)
                if len(races) == 1 and not os.path.lexists(path):
                    os.symlink(external, path)
                return name, path

            with mock.patch.object(convert, "_allocate_asset_path", racing_alloc):
                md_path = convert.convert_docx_to_markdown(docx, out)

            content = Path(md_path).read_text(encoding="utf-8")
            m = re.search(r"!\[\]\((assets/[^)]+)\)", content)
            self.assertIsNotNone(m, content)
            ref = os.path.join(os.path.dirname(md_path), m.group(1))
            self.assertFalse(os.path.islink(ref))
            with open(ref, "rb") as f:
                self.assertEqual(f.read(), png)
            self.assertGreaterEqual(len(races), 2)  # 确实发生了换名重试


class TestSafeRealpathTimeoutPropagation(unittest.TestCase):
    """BUG-039 回归：输出边界检查的路径解析不得吞掉批处理超时信号。

    os.path.realpath（非 strict）内部 except OSError 会把组件 lstat
    抛出的 TimeoutError 当“组件不存在”吞掉并返回字面路径，一次性
    alarm 静默丢失后文档会被误记成功。_safe_realpath 逐段解析，
    TimeoutError 一律上抛；其余行为与 realpath 一致。
    """

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_matches_os_path_realpath(self):
        # 行为一致性：普通/不存在/绝对与相对链接/悬空链接/../ 归一等
        with tempfile.TemporaryDirectory() as tmp:
            real = os.path.join(tmp, "real")
            os.makedirs(os.path.join(real, "sub"))
            link_abs = os.path.join(tmp, "link_abs")
            os.symlink(real, link_abs)
            link_rel = os.path.join(real, "link_rel")
            os.symlink("sub", link_rel)
            dangling = os.path.join(tmp, "dangling")
            os.symlink(os.path.join(tmp, "nope"), dangling)
            for c in (
                real, link_abs, os.path.join(link_abs, "sub"),
                os.path.join(link_rel, "x"), os.path.join(link_abs, "..", "real"),
                dangling, os.path.join(dangling, "tail"),
                os.path.join(tmp, "noexist"), os.path.join(tmp, "noexist", "deeper"),
                real + "/", os.path.join(real, "sub", "..", "sub"), "/",
            ):
                with self.subTest(path=c):
                    self.assertEqual(self.convert._safe_realpath(c), os.path.realpath(c))

    def test_lstat_timeout_error_propagates(self):
        # 判别：旧 realpath 在同场景下吞掉 TimeoutError 返回字面路径
        with mock.patch.object(self.convert.os, "lstat",
                               side_effect=TimeoutError("单文档转换超时")):
            with self.assertRaises(TimeoutError):
                self.convert._safe_realpath("/a/b/c")

    def test_readlink_timeout_error_propagates(self):
        with tempfile.TemporaryDirectory() as tmp:
            link = os.path.join(tmp, "link")
            os.symlink("target", link)
            real_readlink = self.convert.os.readlink

            def readlink_raises_timeout(path):
                # 只对被测链接抛超时，root 探测等其他路径正常
                if str(path).endswith("link"):
                    raise TimeoutError("单文档转换超时")
                return real_readlink(path)

            with mock.patch.object(self.convert.os, "readlink",
                                   side_effect=readlink_raises_timeout):
                with self.assertRaises(TimeoutError):
                    self.convert._safe_realpath(link)


class TestBoundedRead(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")

    def test_bounded_read_rejects_actual_decompression_over_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "input.docx")
            build_minimal_docx(path, extra_entries={"sponge.bin": b"A" * (2 * 1024 * 1024)})
            with zipfile.ZipFile(path, "r") as zf:
                with self.assertRaises(self.convert.DocxSecurityError):
                    self.convert.read_zip_entry_bounded(zf, "sponge.bin", 1024 * 1024)
                # 上限内正常读取
                data = self.convert.read_zip_entry_bounded(zf, "sponge.bin", 4 * 1024 * 1024)
                self.assertEqual(len(data), 2 * 1024 * 1024)


class TestAssetCandidateConflict(unittest.TestCase):
    """BUG-027 回归：hash 消歧候选被错误文件/外部链接/目录占用时，
    必须继续换名并引用内容正确的普通文件。"""

    # 引用了 media 图片的最小 DOCX 构造（与 on_limit 测试的 helper 等价）
    _DOC_NS = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
    )

    @classmethod
    def setUpClass(cls):
        cls.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        cls.image_data = (
            b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR"
            + struct.pack(">II", 2, 2) + b"\x08\x06\x00\x00\x00"
        )
        cls.digest8 = hashlib.sha256(cls.image_data).hexdigest()[:8]

    def _build_docx(self, path):
        drawing = (
            '<w:drawing><wp:inline><wp:extent cx="100" cy="100"/><a:graphic><a:graphicData '
            'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr>'
            '<pic:cNvPr id="1" name="img1"/><pic:cNvPicPr/></pic:nvPicPr><pic:blipFill>'
            '<a:blip r:embed="rId1"/></pic:blipFill><pic:spPr/></pic:pic></a:graphicData>'
            '</a:graphic></wp:inline></w:drawing>'
        )
        document = (
            f'<?xml version="1.0"?><w:document {self._DOC_NS}><w:body>'
            f'<w:p><w:r>{drawing}</w:r></w:p>'
            f'<w:p><w:r><w:t>正文</w:t></w:r></w:p></w:body></w:document>'
        )
        rels = (
            '<?xml version="1.0"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            'Target="media/image1.png"/></Relationships>'
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
            zf.writestr("word/document.xml", document)
            zf.writestr("word/_rels/document.xml.rels", rels)
            zf.writestr("word/media/image1.png", self.image_data)

    def _occupied_candidate_test(self, kind):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = os.path.join(tmp.name, "a.docx")
        self._build_docx(path)
        out = os.path.join(tmp.name, "out")
        assets = os.path.join(out, "a", "assets")
        os.makedirs(assets)
        Path(assets, "image1.png").write_bytes(b"WRONG")  # 自然名已被不同内容占用
        candidate = os.path.join(assets, f"image1_{self.digest8}.png")
        external = Path(tmp.name, "external.bin")
        if kind == "file":
            Path(candidate).write_bytes(b"WRONG2")
        elif kind == "link":
            os.symlink(external, candidate)
        else:
            os.makedirs(candidate)

        md_path = self.convert.convert_docx_to_markdown(path, out)

        content = Path(md_path).read_text(encoding="utf-8")
        match = re.search(r"!\[\]\(assets/([^)]+)\)", content)
        self.assertIsNotNone(match, content)
        target = Path(assets, match.group(1))
        # 输出必须引用内容与输入一致的普通文件，而不是被占用的候选
        self.assertNotEqual(match.group(1), f"image1_{self.digest8}.png")
        self.assertTrue(target.is_file(), match.group(1))
        self.assertFalse(target.is_symlink(), match.group(1))
        self.assertEqual(target.read_bytes(), self.image_data)
        if kind == "link":
            self.assertFalse(external.exists())  # 未跟随链接在外部创建文件

    def test_hash_candidate_occupied_by_wrong_file(self):
        self._occupied_candidate_test("file")

    def test_hash_candidate_occupied_by_external_symlink(self):
        if not hasattr(os, "symlink"):
            self.skipTest("当前平台不支持符号链接")
        # 只捕获“创建符号链接”这一步的环境失败（参照 _symlink_or_skip
        # 的写法）：转换与断言阶段的 OSError 属于真实回归，不得被降级
        # 报成环境跳过（见 TST-013——此前过宽的 except 把
        # _allocate_asset_path 的真实 OSError 也吞成了 skip）
        with tempfile.TemporaryDirectory() as probe_dir:
            try:
                os.symlink(probe_dir, os.path.join(probe_dir, "probe_link"))
            except (OSError, NotImplementedError) as exc:
                self.skipTest(f"当前环境无法创建符号链接: {exc}")
        self._occupied_candidate_test("link")

    def test_hash_candidate_occupied_by_directory(self):
        self._occupied_candidate_test("dir")


class TestAtomicWritePermissions(unittest.TestCase):
    """BUG-028 回归：原子写入不扩大私有权限。"""

    def setUp(self):
        self.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.docx_path = os.path.join(self.tmp.name, "a.docx")
        build_minimal_docx(self.docx_path)

    @unittest.skipUnless(os.name == "posix",
                         "POSIX 权限位断言，Windows 跳过（见 TST-015）")
    def test_new_files_respect_caller_umask(self):
        out = os.path.join(self.tmp.name, "out")
        old_mask = os.umask(0o077)
        try:
            md_path = self.convert.convert_docx_to_markdown(self.docx_path, out)
        finally:
            os.umask(old_mask)

        self.assertEqual(os.stat(md_path).st_mode & 0o777, 0o600)
        sentinel = os.stat(os.path.join(os.path.dirname(md_path), ".converted"))
        self.assertEqual(sentinel.st_mode & 0o777, 0o600)

    @unittest.skipUnless(os.name == "posix",
                         "POSIX 权限位断言，Windows 跳过（见 TST-015）")
    def test_replacement_preserves_existing_file_mode(self):
        out = os.path.join(self.tmp.name, "out")
        md_path = self.convert.convert_docx_to_markdown(self.docx_path, out)
        os.chmod(md_path, 0o600)

        old_mask = os.umask(0o022)  # 默认 umask 下重转也不得扩大已有 0600
        try:
            self.convert.convert_docx_to_markdown(self.docx_path, out)
        finally:
            os.umask(old_mask)

        self.assertEqual(os.stat(md_path).st_mode & 0o777, 0o600)

    def _assert_private_temp_permissions(self, target_name, write):
        """观察真实独占创建后的权限，覆盖空文件被提前打开的窗口。"""
        target = Path(self.tmp.name) / target_name
        target.write_text("旧私有内容", encoding="utf-8")
        target.chmod(0o600)
        real_open = os.open
        observed = []

        def observe_open(path, flags, mode=0o777, **kwargs):
            fd = real_open(path, flags, mode, **kwargs)
            if str(path).endswith(".tmp"):
                st = os.fstat(fd)
                observed.append((st.st_mode & 0o777, st.st_size))
            return fd

        old_mask = os.umask(0o022)
        try:
            with mock.patch("os.open", side_effect=observe_open):
                with mock.patch("os.umask", side_effect=AssertionError("不得查询进程掩码")):
                    write(target)
        finally:
            os.umask(old_mask)

        # 只断言最终0600会漏掉BUG-037：临时文件在创建时也必须私有。
        self.assertEqual(observed, [(0o600, 0)])
        self.assertEqual(target.stat().st_mode & 0o777, 0o600)
        self.assertNotEqual(target.read_text(encoding="utf-8"), "旧私有内容")
        self.assertFalse(list(Path(self.tmp.name).glob("*.tmp")))

    @unittest.skipUnless(os.name == "posix",
                         "POSIX 权限位断言，Windows 跳过（见 TST-015）")
    def test_private_markdown_temp_never_expands_permissions(self):
        self._assert_private_temp_permissions(
            "private.md", lambda path: self.convert._atomic_write_text(
                str(path), "本次私有正文"))

    @unittest.skipUnless(os.name == "posix",
                         "POSIX 权限位断言，Windows 跳过（见 TST-015）")
    def test_private_sentinel_temp_never_expands_permissions(self):
        self._assert_private_temp_permissions(
            ".converted", lambda path: self.convert.write_conversion_sentinel(
                str(path.parent), "private", "private-source-hash"))

    def test_conversion_never_queries_process_umask(self):
        """权限计算不得经 os.umask(0)/恢复查询进程掩码（BUG-034）：
        超时信号落在两步之间会让进程 umask 永久变 0，后续新文件 0666。
        新文件应在独占创建时由内核按 umask 归一权限。"""
        out = os.path.join(self.tmp.name, "out")
        with mock.patch("os.umask",
                        side_effect=AssertionError("不得查询/修改进程 umask")):
            md_path = self.convert.convert_docx_to_markdown(self.docx_path, out)
        self.assertTrue(os.path.isfile(md_path))

    def test_long_cjk_filename_converts_under_255_byte_name_limit(self):
        """临时文件名不得依赖目标名长度（BUG-036）：80 汉字 stem 的 md 名
        243 字节；ext4 单名按 255 字节计，旧实现按目标名拼临时前缀达
        257 字节会 ENAMETOOLONG。APFS 按字符数计长、本机直测不触发，
        故用 os.open 包装模拟 ext4 字节上限。"""
        stem = "测" * 80
        docx_path = os.path.join(self.tmp.name, f"{stem}.docx")
        build_minimal_docx(docx_path)
        out = os.path.join(self.tmp.name, "out")
        real_open = os.open

        def ext4_open(path, flags, mode=0o777, *, dir_fd=None):
            if len(os.path.basename(str(path)).encode()) > 255:
                raise OSError(errno.ENAMETOOLONG, "File name too long", str(path))
            return real_open(path, flags, mode, dir_fd=dir_fd)

        with mock.patch("os.open", side_effect=ext4_open):
            md_path = self.convert.convert_docx_to_markdown(docx_path, out)

        self.assertEqual(os.path.basename(md_path), f"{stem}.md")
        self.assertTrue(os.path.isfile(md_path))


class TestAssetCompareTimeoutNotSwallowed(unittest.TestCase):
    """BUG-025 验收补充：资源比较助手不得吞掉批处理超时信号。"""

    def setUp(self):
        self.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        self.batch = load_module("batch_convert_module", SCRIPTS_DIR / "batch_convert.py")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.src = os.path.join(self.tmp.name, "src")
        self.out = os.path.join(self.tmp.name, "out")
        os.makedirs(self.src)

    def test_same_file_content_reraises_timeout_error(self):
        Path(self.tmp.name, "existing.bin").write_bytes(b"x")
        with mock.patch("os.path.getsize", side_effect=TimeoutError("单文档转换超时")):
            with self.assertRaises(TimeoutError):
                self.convert._same_file_content(
                    os.path.join(self.tmp.name, "existing.bin"), b"data")

    def test_atomic_write_text_reraises_timeout_error_on_target_lstat(self):
        """目标权限查询（lstat）上的 TimeoutError 不得被 OSError 捕获吞掉（BUG-030）。"""
        target_dir = os.path.join(self.tmp.name, "w")
        os.makedirs(target_dir)
        with mock.patch("os.lstat", side_effect=TimeoutError("单文档转换超时")):
            with self.assertRaises(TimeoutError):
                self.convert._atomic_write_text(
                    os.path.join(target_dir, "out.md"), "x")
        self.assertEqual(os.listdir(target_dir), [])  # 临时文件已清理

    @unittest.skipUnless(hasattr(signal, "SIGALRM"),
                         "平台无 SIGALRM（Windows），超时路径自动跳过")
    def test_batch_timeout_during_asset_compare_counts_failure(self):
        # image1.png 与 image1.jpeg（PNG 魔数，扩展名修正后同名、内容不同）
        # 在同一次转换内触发内容比较；比较期间 SIGALRM 触发必须计失败。
        path = os.path.join(self.src, "a.docx")
        doc = ('<?xml version="1.0"?><w:document '
               'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               '<w:body><w:p><w:r><w:t>x</w:t></w:r></w:p></w:body></w:document>')
        png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(">II", 2, 2) \
            + b"\x08\x06\x00\x00\x00"
        other = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(">II", 3, 3) \
            + b"\x08\x06\x00\x00\x00"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
            zf.writestr("word/document.xml", doc)
            zf.writestr("word/media/image1.png", png)
            zf.writestr("word/media/image1.jpeg", other)

        with mock.patch("os.path.getsize", side_effect=lambda p: time.sleep(2)):
            with self.assertLogs("batch_convert_module", level="INFO"):
                summary = self.batch.batch_convert(self.src, self.out, timeout=1)

        self.assertEqual(summary, {"success": 0, "skipped": 0, "failed": 1})
        self.assertFalse(os.path.exists(os.path.join(self.out, "a")))  # 半成品被清理

    @unittest.skipUnless(hasattr(signal, "SIGALRM"),
                         "平台无 SIGALRM（Windows），超时路径自动跳过")
    def test_batch_timeout_during_permission_lstat_counts_failure(self):
        """权限查询里的 lstat 不得吞掉 SIGALRM（BUG-030）。"""
        path = os.path.join(self.src, "b.docx")
        build_minimal_docx(path)
        real_lstat = os.lstat

        def slow_lstat(target, *args, **kwargs):
            if "_atomic_write_text" in "".join(__import__("traceback").format_stack()) \
                    and str(target).endswith(".md"):
                time.sleep(2)
            return real_lstat(target, *args, **kwargs)

        with mock.patch("os.lstat", side_effect=slow_lstat):
            summary = self.batch.batch_convert(self.src, self.out, timeout=1)

        self.assertEqual(summary, {"success": 0, "skipped": 0, "failed": 1})
        self.assertFalse(os.path.exists(os.path.join(self.out, "b")))

    @unittest.skipUnless(hasattr(signal, "SIGALRM"),
                         "平台无 SIGALRM（Windows），超时路径自动跳过")
    def test_batch_timeout_during_asset_path_lstat_counts_failure(self):
        """图片候选路径检查不得吞掉 SIGALRM（BUG-035）：
        os.path.lexists/islink/isfile 内部捕获 OSError（含其子类
        TimeoutError），必须改为直接 lstat 并显式重抛。"""
        path = os.path.join(self.src, "c.docx")
        png = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" \
            + struct.pack(">II", 2, 2) + b"\x08\x06\x00\x00\x00"
        build_minimal_docx(path, media_entries={"image1.png": png})
        real_lstat = os.lstat

        def slow_lstat(target, *args, **kwargs):
            # 仅延迟 assets 候选名的 lstat，输出目录检查不受影响
            if os.path.basename(str(target)).startswith("image") \
                    and str(os.path.dirname(str(target))).endswith("assets"):
                time.sleep(2)
            return real_lstat(target, *args, **kwargs)

        with mock.patch("os.lstat", side_effect=slow_lstat):
            with self.assertLogs("batch_convert_module", level="INFO"):
                summary = self.batch.batch_convert(self.src, self.out, timeout=1)

        self.assertEqual(summary, {"success": 0, "skipped": 0, "failed": 1})
        self.assertFalse(os.path.exists(os.path.join(self.out, "c")))  # 半成品被清理


class TestOutputSymlinkEscape(unittest.TestCase):
    """issue #4 回归（[[BUG-019]]/[[BUG-023]]/[[BUG-024]]）：
    输出路径上预置的符号链接不得把删除或写入导向输出根目录之外。"""

    def setUp(self):
        self.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _symlink_or_skip(self, target, link):
        try:
            os.symlink(target, link)
        except (OSError, NotImplementedError) as exc:
            # Windows 有 os.symlink 但默认无创建特权（WinError 1314）
            self.skipTest(f"当前环境无法创建符号链接: {exc}")

    def _build_plain_docx(self, name="evil.docx"):
        path = os.path.join(self.tmp.name, name)
        build_minimal_docx(path)
        return path

    def test_output_subdir_symlink_rejected_external_untouched(self):
        # output/evil -> victim：转换必须整体拒绝，victim 内容零改动
        out = os.path.join(self.tmp.name, "out")
        victim = os.path.join(self.tmp.name, "victim")
        os.makedirs(os.path.join(victim, "assets"))
        keep = Path(victim, "assets", "keep.bin")
        keep.write_bytes(b"KEEP")
        os.makedirs(out)
        self._symlink_or_skip(victim, os.path.join(out, "evil"))

        with self.assertRaisesRegex(ValueError, "不允许符号链接"):
            self.convert.convert_docx_to_markdown(self._build_plain_docx(), out)

        self.assertEqual(keep.read_bytes(), b"KEEP")
        self.assertEqual(os.listdir(victim), ["assets"])
        self.assertFalse(os.path.exists(os.path.join(victim, "evil.md")))
        self.assertTrue(os.path.islink(os.path.join(out, "evil")))  # 链接本身保留

    def test_output_md_symlink_replaced_without_following(self):
        # 普通目录内预置 evil.md -> 外部 keep.txt：转换成功，但外部内容
        # 零改动，输出位置的链接被替换为本目录普通文件
        out_doc = os.path.join(self.tmp.name, "out", "evil")
        os.makedirs(out_doc)
        keep = Path(self.tmp.name, "keep.txt")
        keep.write_text("KEEP", encoding="utf-8")
        self._symlink_or_skip(keep, os.path.join(out_doc, "evil.md"))

        md_path = self.convert.convert_docx_to_markdown(
            self._build_plain_docx(), os.path.join(self.tmp.name, "out"))

        self.assertEqual(keep.read_text(encoding="utf-8"), "KEEP")
        self.assertFalse(os.path.islink(md_path))
        self.assertIn("测试正文", Path(md_path).read_text(encoding="utf-8"))

    def test_sentinel_fixed_name_tmp_symlink_not_followed(self):
        # 预置固定名 .converted.tmp -> 外部 keep.txt：sentinel 写入不得
        # 覆盖外部内容，也不得把该链接 rename 成 .converted
        out_doc = os.path.join(self.tmp.name, "out", "evil")
        os.makedirs(out_doc)
        keep = Path(self.tmp.name, "keep2.txt")
        keep.write_text("KEEP", encoding="utf-8")
        self._symlink_or_skip(keep, os.path.join(out_doc, ".converted.tmp"))

        self.convert.convert_docx_to_markdown(
            self._build_plain_docx(), os.path.join(self.tmp.name, "out"))

        self.assertEqual(keep.read_text(encoding="utf-8"), "KEEP")
        self.assertFalse(os.path.islink(os.path.join(out_doc, ".converted")))
        self.assertIsNotNone(self.convert.read_conversion_sentinel(out_doc))

    def test_assets_file_symlink_not_followed(self):
        # 预置悬空 assets/image1.png -> 外部（尚不存在的）keep3.bin：
        # 写入不得跟随链接在外部创建文件，图片改用 hash 后缀名落盘
        path = os.path.join(self.tmp.name, "img.docx")
        build_minimal_docx(
            path, media_entries={"image1.png": b"\x89PNG\r\n\x1a\n" + b"\0" * 32})
        assets = os.path.join(self.tmp.name, "out", "img", "assets")
        os.makedirs(assets)
        external = Path(self.tmp.name, "keep3.bin")
        self._symlink_or_skip(external, os.path.join(assets, "image1.png"))

        self.convert.convert_docx_to_markdown(path, os.path.join(self.tmp.name, "out"))

        self.assertFalse(external.exists())  # 未跟随链接在外部创建文件
        # 悬空链接被清理，实际图片以 hash 后缀名落盘
        extracted = os.listdir(assets)
        self.assertEqual(len(extracted), 1, extracted)
        self.assertTrue(extracted[0].startswith("image1_"), extracted)
        self.assertFalse(os.path.islink(os.path.join(assets, extracted[0])))


def _build_xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["col1", "col2"])
    ws.append(["1", "2"])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class TestExcelTimeoutNotSwallowed(unittest.TestCase):
    """BUG-025 回归：批处理 SIGALRM 超时不得被 Excel 降级捕获吞掉。"""

    def setUp(self):
        self.convert = load_module("convert_docx_module", SCRIPTS_DIR / "convert_docx.py")
        self.batch = load_module("batch_convert_module", SCRIPTS_DIR / "batch_convert.py")
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.src = os.path.join(self.tmp.name, "src")
        self.out = os.path.join(self.tmp.name, "out")
        os.makedirs(self.src)
        self.docx_path = os.path.join(self.src, "input.docx")
        build_minimal_docx(self.docx_path)
        with zipfile.ZipFile(self.docx_path, "a", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("word/embeddings/sheet1.xlsx", _build_xlsx_bytes())
        self.target_dir = os.path.join(self.out, "input")

    def test_excel_to_markdown_reraises_timeout_error(self):
        with mock.patch("openpyxl.load_workbook", side_effect=TimeoutError("单文档转换超时")):
            with self.assertRaises(TimeoutError):
                self.convert.excel_to_markdown(_build_xlsx_bytes())

    def test_textbox_extraction_reraises_timeout(self):
        # 其他降级捕获点同样不得吞掉超时信号
        with mock.patch.object(
            self.convert, "_safe_xml_fromstring", side_effect=TimeoutError("单文档转换超时")
        ):
            with self.assertRaises(TimeoutError):
                self.convert.extract_textbox_content(self.docx_path)

    @unittest.skipUnless(hasattr(signal, "SIGALRM"),
                         "平台无 SIGALRM（Windows），超时路径自动跳过")
    def test_batch_timeout_inside_excel_conversion_counts_failure(self):
        # 超时在真实转换内部（Excel 解析阶段）触发：必须计失败、清理输出、
        # 不写 sentinel，而不是被降级为“无表格的成功转换”
        def slow_load(*args, **kwargs):
            time.sleep(2)

        with mock.patch("openpyxl.load_workbook", side_effect=slow_load):
            with self.assertLogs("batch_convert_module", level="INFO") as cm:
                summary = self.batch.batch_convert(self.src, self.out, timeout=1)

        self.assertEqual(summary, {"success": 0, "skipped": 0, "failed": 1})
        self.assertTrue(any("超时" in line for line in cm.output))
        self.assertFalse(os.path.exists(self.target_dir))  # 半成品被清理


if __name__ == "__main__":
    unittest.main()
